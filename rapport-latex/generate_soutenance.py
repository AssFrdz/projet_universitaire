from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
FIGURES = ROOT / "rapport" / "figures"
OUTPUT_PPTX = ROOT / "soutenance-new.pptx"


NAVY = RGBColor(26, 42, 79)
BLUE = RGBColor(52, 89, 149)
LIGHT_BLUE = RGBColor(232, 239, 250)
TEXT = RGBColor(32, 37, 43)
MUTED = RGBColor(95, 103, 114)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(46, 125, 50)
RED = RGBColor(181, 54, 46)


def set_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_top_band(slide, title, subtitle=None):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.85))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.color.rgb = NAVY

    title_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.10), Inches(10.8), Inches(0.34))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.name = "Aptos Display"
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = WHITE

    if subtitle:
        p2 = tf.add_paragraph()
        p2.level = 0
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.name = "Aptos"
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(220, 227, 243)


def add_footer(slide, page_no):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(7.06), Inches(12.43), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT_BLUE
    line.line.color.rgb = LIGHT_BLUE

    left_box = slide.shapes.add_textbox(Inches(0.45), Inches(7.10), Inches(6.0), Inches(0.22))
    tf_left = left_box.text_frame
    tf_left.clear()
    p_left = tf_left.paragraphs[0]
    p_left.alignment = PP_ALIGN.LEFT
    r_left = p_left.add_run()
    r_left.text = "Projet universitaire MIAGE - Marmote"
    r_left.font.name = "Aptos"
    r_left.font.size = Pt(9)
    r_left.font.color.rgb = MUTED

    right_box = slide.shapes.add_textbox(Inches(12.2), Inches(7.08), Inches(0.5), Inches(0.22))
    tf_right = right_box.text_frame
    tf_right.clear()
    p_right = tf_right.paragraphs[0]
    p_right.alignment = PP_ALIGN.RIGHT
    r_right = p_right.add_run()
    r_right.text = str(page_no)
    r_right.font.name = "Aptos"
    r_right.font.size = Pt(9)
    r_right.font.color.rgb = MUTED


def add_bullets(slide, items, left, top, width, height, font_size=22, color=TEXT, bullet_color=BLUE):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.clear()

    for idx, item in enumerate(items):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        p.line_spacing = 1.1
        r = p.add_run()
        r.text = "- " + item
        r.font.name = "Aptos"
        r.font.size = Pt(font_size)
        r.font.color.rgb = color


def add_paragraph_block(slide, lines, left, top, width, height, font_size=19, color=TEXT, bold_first=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        p.line_spacing = 1.05
        r = p.add_run()
        r.text = line
        r.font.name = "Aptos"
        r.font.size = Pt(font_size)
        r.font.color.rgb = color
        if bold_first and idx == 0:
            r.font.bold = True


def add_label_box(slide, title, body, left, top, width, height, fill_color=LIGHT_BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color

    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True

    p1 = tf.paragraphs[0]
    p1.space_after = Pt(6)
    r1 = p1.add_run()
    r1.text = title
    r1.font.name = "Aptos Display"
    r1.font.size = Pt(18)
    r1.font.bold = True
    r1.font.color.rgb = NAVY

    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = body
    r2.font.name = "Aptos"
    r2.font.size = Pt(16)
    r2.font.color.rgb = TEXT


def add_picture(slide, image_name, left, top, width=None, height=None):
    path = FIGURES / image_name
    if not path.exists():
        raise FileNotFoundError(path)
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def add_caption(slide, text, left, top, width, size=11):
    box = slide.shapes.add_textbox(left, top, width, Inches(0.28))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = MUTED


def fit_pair_heights(pic1, pic2, target_top, max_height):
    ratio1 = pic1.width / pic1.height
    ratio2 = pic2.width / pic2.height
    pic1.height = max_height
    pic1.width = int(max_height * ratio1)
    pic2.height = max_height
    pic2.width = int(max_height * ratio2)
    pic1.top = target_top
    pic2.top = target_top


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    # Slide 1
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    banner.fill.solid()
    banner.fill.fore_color.rgb = WHITE
    banner.line.color.rgb = WHITE
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.05))
    top.fill.solid()
    top.fill.fore_color.rgb = NAVY
    top.line.color.rgb = NAVY

    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(1.45), Inches(11.9), Inches(1.6))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "Stabilisation du workflow documentaire de Marmote"
    r.font.name = "Aptos Display"
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = NAVY
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "et integration de notebooks C++"
    r2.font.name = "Aptos Display"
    r2.font.size = Pt(28)
    r2.font.bold = True
    r2.font.color.rgb = NAVY

    sub = slide.shapes.add_textbox(Inches(0.78), Inches(3.25), Inches(7.4), Inches(1.2))
    stf = sub.text_frame
    stf.clear()
    for text, size, bold, color in [
        ("Assia FERRADJ", 21, True, TEXT),
        ("Projet universitaire - Licence MIAGE", 16, False, MUTED),
        ("Encadrement : Emmanuel Hyon", 16, False, MUTED),
    ]:
        p = stf.paragraphs[0] if not stf.text else stf.add_paragraph()
        r = p.add_run()
        r.text = text
        r.font.name = "Aptos"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        p.space_after = Pt(6)

    add_paragraph_block(
        slide,
        [
            "Objectif : fiabiliser la chaine de generation documentaire et integrer",
            "des notebooks C++ executables et publiables dans Sphinx."
        ],
        Inches(0.8), Inches(4.55), Inches(8.2), Inches(0.9), font_size=18, color=BLUE
    )
    add_picture(slide, "cpp_api_page_real.png", Inches(9.4), Inches(1.55), width=Inches(3.1))

    # Slide 2
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_top_band(slide, "Contexte et problematique")
    add_bullets(
        slide,
        [
            "Marmote est une librairie de modelisation markovienne avec une API C++ et une API Python.",
            "La documentation est publiee avec Sphinx a partir de notebooks, d'exemples C++ et d'archives telechargeables.",
            "Le depot initial etait fonctionnel, mais fragmenté : scripts multiples, chemins fragiles, contenus C++ encore peu integrés.",
            "Problematique : comment ajouter des notebooks C++ fiables et maintenables sans casser le workflow documentaire existant ?",
        ],
        Inches(0.7), Inches(1.25), Inches(7.0), Inches(4.9), font_size=20
    )
    add_picture(slide, "cpp_api_page_real.png", Inches(8.25), Inches(1.55), width=Inches(4.45))
    add_caption(slide, "Exemple de page de documentation generee", Inches(8.05), Inches(6.1), Inches(4.8))
    add_footer(slide, 2)

    # Slide 3
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_top_band(slide, "Objectifs et demarche")
    add_label_box(
        slide,
        "Objectifs",
        "Stabiliser les scripts, integrer des notebooks C++, generer des exports propres, et rendre le depot livrable.",
        Inches(0.7), Inches(1.25), Inches(5.9), Inches(1.2)
    )
    add_label_box(
        slide,
        "Demarche",
        "Stabiliser d'abord l'existant, puis etendre progressivement, en validant chaque ajout technique a petite echelle.",
        Inches(6.75), Inches(1.25), Inches(5.85), Inches(1.2)
    )
    add_bullets(
        slide,
        [
            "S'appuyer sur deux references : notebooks Python pour la pedagogie, exemples xpl/ pour la verite technique.",
            "Rendre le workflow reproductible : dry run, redirection du repertoire courant, point d'entree unique `total.py`.",
            "Traiter la soutenance comme un livrable : clarté du depot, sorties propres, documentation developpeur, archives coherentes.",
        ],
        Inches(0.8), Inches(2.95), Inches(11.7), Inches(2.7), font_size=21
    )
    add_footer(slide, 3)

    # Slide 4
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_top_band(slide, "Evolution du workflow")
    add_paragraph_block(
        slide,
        ["Passage d'un flux fragmente a un flux unifie, pilote par `total.py`."],
        Inches(0.78), Inches(1.0), Inches(8.6), Inches(0.4), font_size=18, color=MUTED
    )
    add_picture(slide, "workflow_base.png", Inches(0.75), Inches(1.5), width=Inches(5.95))
    add_picture(slide, "workflow_cible.png", Inches(6.7), Inches(1.5), width=Inches(5.9))
    add_caption(slide, "Workflow initial", Inches(0.9), Inches(5.85), Inches(5.5))
    add_caption(slide, "Workflow final", Inches(6.95), Inches(5.85), Inches(5.3))
    add_footer(slide, 4)

    # Slide 5
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_top_band(slide, "Organisation et environnements de travail")
    add_bullets(
        slide,
        [
            "Windows pour les scripts, Sphinx et la gestion du depot.",
            "WSL pour l'execution des notebooks C++ via Xeus-cling.",
            "Environnement `marmote-use` pour la documentation et les builds Python.",
            "Environnement `xeus-cpp-env` pour JupyterLab, `xcpp17` et Marmote cote C++.",
        ],
        Inches(0.75), Inches(1.35), Inches(6.0), Inches(4.7), font_size=21
    )
    add_picture(slide, "jupyter_server_wsl.png", Inches(7.15), Inches(1.45), width=Inches(2.55))
    add_picture(slide, "kernelspec_list.png", Inches(9.9), Inches(1.45), width=Inches(2.55))
    add_caption(slide, "Serveur Jupyter lance dans WSL", Inches(7.1), Inches(5.35), Inches(2.7))
    add_caption(slide, "Verification du kernel C++", Inches(9.85), Inches(5.35), Inches(2.7))
    add_footer(slide, 5)

    # Slide 6
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_top_band(slide, "Stabilisation du workflow existant")
    add_bullets(
        slide,
        [
            "Correction des incoherences d'arborescence et des destinations d'archives.",
            "Realignement des correspondances entre dossiers source, executables et parametres d'execution.",
            "Ajout de `dry_run`, de controles d'existence et d'une redirection automatique vers `doc/html/source`.",
            "Consolidation d'un point d'entree unique : `scripts/total.py`.",
        ],
        Inches(0.75), Inches(1.35), Inches(6.25), Inches(4.9), font_size=20
    )
    add_picture(slide, "redirection_test.png", Inches(7.3), Inches(1.55), width=Inches(5.2))
    add_caption(slide, "Exemple de redirection de contexte vers le bon dossier de travail", Inches(7.0), Inches(5.75), Inches(5.7))
    add_footer(slide, 6)

    # Slide 7
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_top_band(slide, "Creation des notebooks C++")
    add_bullets(
        slide,
        [
            "Creation d'un notebook C++ dans WSL avec le kernel `xcpp17`.",
            "Ajout de cellules techniques `#pragma cling` pour charger Marmote, puis masquage de ces cellules dans la documentation.",
            "Reecriture des lecons pour suivre la structure pedagogique des notebooks Python.",
            "Correction de plusieurs incompatibilites d'API au fil des essais dans Cling.",
        ],
        Inches(0.72), Inches(1.28), Inches(5.75), Inches(4.9), font_size=19
    )
    add_picture(slide, "select_kernel_cpp17.png", Inches(6.75), Inches(1.45), width=Inches(2.2))
    add_picture(slide, "marmote_cpp_notebook_setup.png", Inches(9.15), Inches(1.45), width=Inches(3.55))
    add_picture(slide, "marmote_cpp_notebook_output.png", Inches(6.75), Inches(4.15), width=Inches(5.95))
    add_caption(slide, "Selection du kernel puis configuration Marmote et premier resultat", Inches(6.55), Inches(6.08), Inches(6.2))
    add_footer(slide, 7)

    # Slide 8
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_top_band(slide, "Exports C++, publication et archives")
    add_bullets(
        slide,
        [
            "Mise en place de `creating_cppFiles(...)` pour exporter les notebooks C++ en `.cpp` via `nbconvert.ScriptExporter`.",
            "Suppression des directives Xeus-cling dans les exports et exclusion des cellules cachees.",
            "Generation automatique d'un `CMakeLists.txt` dans l'archive des notebooks C++.",
            "Mise a jour des index Sphinx pour exposer les nouvelles lecons `cpptutos/`.",
        ],
        Inches(0.72), Inches(1.22), Inches(6.05), Inches(4.95), font_size=19
    )
    add_picture(slide, "cpp_examples_old.png", Inches(7.05), Inches(1.35), width=Inches(2.65))
    add_picture(slide, "cpp_examples_new.png", Inches(9.95), Inches(1.35), width=Inches(2.65))
    add_picture(slide, "tag_cells.png", Inches(7.1), Inches(4.0), width=Inches(5.55))
    add_caption(slide, "Refonte de l'index C++ et gestion des cellules techniques", Inches(7.0), Inches(6.05), Inches(5.8))
    add_footer(slide, 8)

    # Slide 9
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_top_band(slide, "Difficultes rencontrees et solutions")
    add_label_box(slide, "Arborescence fragile", "Reconstruction des chemins, mappings et dossiers de destination.", Inches(0.7), Inches(1.25), Inches(4.0), Inches(1.0))
    add_label_box(slide, "Xeus indisponible sous Windows", "Bascule vers WSL pour les notebooks C++ et maintien de Sphinx sous Windows.", Inches(4.85), Inches(1.25), Inches(4.0), Inches(1.0))
    add_label_box(slide, "Incoherences API Marmote", "Appui systematique sur `xpl/` et corrections ciblees dans les notebooks.", Inches(9.0), Inches(1.25), Inches(3.6), Inches(1.0))
    add_label_box(slide, "Probleme CMake", "Clarification du besoin : indiquer `CMAKE_PREFIX_PATH` ou `marmote_DIR`.", Inches(0.7), Inches(2.55), Inches(5.5), Inches(0.95), fill_color=RGBColor(246, 240, 230))
    add_label_box(slide, "Sphinx et notebooks C++", "Build Windows en mode documentaire avec `nbsphinx_execute=never` si le kernel `xcpp17` reste uniquement dans WSL.", Inches(6.45), Inches(2.55), Inches(6.2), Inches(0.95), fill_color=RGBColor(246, 240, 230))
    add_picture(slide, "cmake_error_find_package.png", Inches(0.95), Inches(3.85), width=Inches(5.6))
    add_picture(slide, "cmake_success_build.png", Inches(6.8), Inches(3.85), width=Inches(5.6))
    add_caption(slide, "Echec initial puis configuration reussie apres indication du bon chemin Marmote", Inches(1.0), Inches(6.02), Inches(11.3))
    add_footer(slide, 9)

    # Slide 10
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_top_band(slide, "Resultats obtenus")
    add_bullets(
        slide,
        [
            "Un workflow plus lisible : compilation, generation d'archives et publication Sphinx mieux articulees.",
            "Des notebooks C++ integrés dans `cpptutos/`, alignes sur la pedagogie des notebooks Python.",
            "Des exports `.cpp` propres, reexploitables, avec un `CMakeLists.txt` distribue dans l'archive.",
            "Un depot plus livrable : README developpeur, commentaires homogenises et documentation regeneratee.",
        ],
        Inches(0.75), Inches(1.35), Inches(6.0), Inches(4.9), font_size=20
    )
    add_picture(slide, "cpp_api_page_real.png", Inches(7.25), Inches(1.5), width=Inches(5.0))
    add_picture(slide, "colab_install_check.png", Inches(8.15), Inches(4.55), width=Inches(3.2))
    add_caption(slide, "Documentation finale et adaptation des notebooks Python pour Colab", Inches(7.0), Inches(6.05), Inches(5.6))
    add_footer(slide, 10)

    # Slide 11
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_top_band(slide, "Conclusion et perspectives")
    add_label_box(slide, "Bilan", "Le projet a surtout consisté a stabiliser un workflow documentaire complet avant de l'etendre.", Inches(0.78), Inches(1.25), Inches(5.7), Inches(1.0))
    add_bullets(
        slide,
        [
            "Montee en competence sur C++, notebooks Jupyter, Sphinx, Conda, WSL et CMake.",
            "Approche de maintenance evolutive : comprendre l'existant, corriger sans casser, puis etendre.",
            "Perspectives : CI de build, tests automatiques, fichiers d'environnement, et nettoyage des warnings Sphinx restants.",
        ],
        Inches(0.8), Inches(2.55), Inches(7.3), Inches(3.0), font_size=20
    )
    add_picture(slide, "workflow_cible.png", Inches(8.55), Inches(1.55), width=Inches(4.0))
    add_footer(slide, 11)

    # Slide 12
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.05))
    top.fill.solid()
    top.fill.fore_color.rgb = NAVY
    top.line.color.rgb = NAVY

    box = slide.shapes.add_textbox(Inches(0.85), Inches(1.7), Inches(8.2), Inches(1.2))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Merci pour votre attention"
    r.font.name = "Aptos Display"
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = NAVY
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Questions ?"
    r2.font.name = "Aptos Display"
    r2.font.size = Pt(26)
    r2.font.bold = True
    r2.font.color.rgb = BLUE

    add_paragraph_block(
        slide,
        [
            "Assia FERRADJ",
            "assia.ferradj@parisnanterre.fr",
            "assfrdz.github.io/portfolio-assia-ferradj-v2",
            "linkedin.com/in/assia-ferradj-895758296",
        ],
        Inches(0.9), Inches(3.35), Inches(6.9), Inches(2.1), font_size=18
    )
    add_picture(slide, "logo_nanterre.png", Inches(9.15), Inches(1.8), width=Inches(2.8))
    add_picture(slide, "logo_Paris_Nanterre_couleur_RVB.png", Inches(7.7), Inches(4.15), width=Inches(4.0))

    prs.save(str(OUTPUT_PPTX))


if __name__ == "__main__":
    build_presentation()
